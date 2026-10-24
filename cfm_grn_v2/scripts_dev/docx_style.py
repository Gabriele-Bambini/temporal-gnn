"""Reference document for pandoc and post-processing of the generated Word report (A4, fonts, tables, footer)."""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def _font(style, name, size=None, bold=None, italic=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def make_reference_docx(path: Path) -> Path:
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(str(path))
    for sec in d.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    st = d.styles

    def get(name):
        try:
            return st[name]
        except KeyError:
            return None

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        s = get(name)
        if s is not None:
            _font(s, "Calibri", 10.5)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.line_spacing = 1.12
    for name, size, before in (("Heading 1", 15, 16), ("Heading 2", 12.5, 11), ("Heading 3", 11, 8)):
        s = get(name)
        if s is not None:
            _font(s, "Calibri", size, bold=True, color=NAVY)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.keep_with_next = True
    s = get("Title")
    if s is not None:
        _font(s, "Calibri", 21, bold=True, color=NAVY)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.space_after = Pt(6)
    s = get("Subtitle")
    if s is not None:
        _font(s, "Calibri", 13, bold=False, italic=False, color=GREY)
        s.paragraph_format.space_after = Pt(10)
    for name in ("Author", "Date"):
        s = get(name)
        if s is not None:
            _font(s, "Calibri", 11, color=GREY)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name in ("Image Caption", "Table Caption", "Caption"):
        s = get(name)
        if s is not None:
            _font(s, "Calibri", 9, italic=True, color=GREY)
            s.paragraph_format.space_after = Pt(10)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if name == "Table Caption" else WD_ALIGN_PARAGRAPH.CENTER
            if name == "Table Caption":
                s.paragraph_format.keep_with_next = True
    s = get("Source Code")
    if s is not None:
        _font(s, "Consolas", 8)
    s = get("Verbatim Char")
    if s is not None:
        _font(s, "Consolas", 9.5)
    s = get("TOC Heading")
    if s is not None:
        _font(s, "Calibri", 13, bold=True, color=NAVY)
    s = get("Table")
    if s is not None:
        tbl_pr = s.element.find(qn("w:tblPr"))
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            s.element.append(tbl_pr)
        old = tbl_pr.find(qn("w:tblBorders"))
        if old is not None:
            tbl_pr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for edge, val, sz, color in (("top", "single", 8, "1F4E79"), ("bottom", "single", 8, "1F4E79"),
                                     ("insideH", "single", 4, "C9D3DF"), ("left", "nil", 0, "auto"),
                                     ("right", "nil", 0, "auto"), ("insideV", "nil", 0, "auto")):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), val)
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            borders.append(e)
        tbl_pr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 40), ("bottom", 40), ("left", 80), ("right", 80)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(w))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tbl_pr.append(mar)
    d.save(str(path))
    return path


def _page_field(paragraph, size=8):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = GREY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for e in (begin, instr, sep, txt, end):
        run._r.append(e)


def polish_docx(path: Path, footer_text: str) -> None:
    d = Document(str(path))
    for t in d.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        n_rows = len(t.rows)
        for ri, row in enumerate(t.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for p in cell.paragraphs:                 # keep every table on one page
                    p.paragraph_format.keep_with_next = ri < n_rows - 1
                    p.paragraph_format.keep_together = True
            for cell in row.cells:
                if ri == 0:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "E8EEF7")
                    tc_pr.append(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if ri == 0:
                            r.font.bold = True
    for sec in d.sections:
        f = sec.footer
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(footer_text + "   ·   page ")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        _page_field(p)
    d.save(str(path))
